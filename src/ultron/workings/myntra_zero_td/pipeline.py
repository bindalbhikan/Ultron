"""Myntra zero TD (10 AM, 2 PM, 5 PM, Monday to Saturday): which styles of the MYNTRA ZERO TD
UPDATE SHEET are live on the buyer side (www.myntra.com) with stock and selling at MRP, i.e. with
no trade discount on.

The sheet (Drive AGENT WORK/MYNTRA/MYNTRA ZERO TD UPDATE SHEET.xlsx) has one tab per account and month,
"MAIN-SEP-26", "DBCC-SEP-26", "FMCC-SEP-26" (ARTICLE TYPE, BRAND, STYLE ID, MY SKU, FINAL SKU, COST,
LIQ STYLES); the last tab of each account is the one read. Every style is looked up on the buyer
side (spider.py). A style is
  ZERO TD       every size in stock sells at MRP
  PART ZERO TD  some sizes in stock sell at MRP, others with a discount
  DISCOUNTED    every size in stock sells below MRP
  OUT OF STOCK  no size can be bought
  NOT FOUND     not on Myntra
Telegram gets one message (per account the zero TD styles and their stock, the biggest ones) and
one workbook: a tab per account with its ZERO TD and PART ZERO TD styles, SUMMARY, and NOT CHECKED
when some styles had no answer. The workbook is also kept on Drive:
AGENT WORK/MYNTRA/MYNTRA ZERO TD SHEETS/OCT-26/26-10-26/ZERO TD 26-10-26 10.00 AM.xlsx.
"""
from __future__ import annotations

import argparse
import logging
from datetime import datetime
from pathlib import Path

import pandas as pd

from ..core import telegram
from ..core.app import run_bot
from ..core.config import csv_list, env, num, output_dir

log = logging.getLogger("bindal.myntra_zero_td")

AGENT_WORK = "1uO-aOha_QKkI3f19EA7sweYPgkjMukZI"   # Drive BINDAL EXPORTS/AGENT WORK
AREA = "MYNTRA"   # the agent's folder is kept by portal: AGENT WORK/MYNTRA (since 29 Oct 2026)
SOURCE = "MYNTRA ZERO TD UPDATE SHEET.xlsx"
FOLDER = "MYNTRA ZERO TD SHEETS"   # AGENT WORK/MYNTRA/<FOLDER>/<OCT-26>/<26-10-26>/ZERO TD 26-10-26 10.00 AM.xlsx
ACCOUNTS = ["MAIN", "DBCC", "FMCC"]
SHEET_COLUMNS = ["ARTICLE TYPE", "BRAND", "STYLE ID", "MY SKU", "FINAL SKU", "COST", "LIQ STYLES"]
CHECK_COLUMNS = ["STATUS", "MRP", "SELLING PRICE", "SIZES AT MRP", "STOCK AT MRP", "TOTAL STOCK",
                 "SELLER", "LINK"]
FLAGGED = ("ZERO TD", "PART ZERO TD")
ALL_TAB = "ALL ACCOUNTS"
TOP = 5               # styles named per account in the message
MAX_UNCHECKED = 0.05  # more styles than this without an answer fails the run


def read_source(path: Path, accounts: list[str]) -> dict[str, pd.DataFrame]:
    """{account: its tab's rows, one per style id}; the last tab of each account wins."""
    tabs = pd.read_excel(path, sheet_name=None, dtype=object)
    out: dict[str, pd.DataFrame] = {}
    for title, df in tabs.items():
        account = str(title).split("-")[0].strip().upper()
        if account not in accounts:
            continue
        df.columns = [str(c).strip().upper() for c in df.columns]
        if "STYLE ID" not in df.columns:
            log.warning("Tab %s has no STYLE ID column: skipped", title)
            continue
        ids = pd.to_numeric(df["STYLE ID"], errors="coerce")
        df = df[ids.notna()].assign(**{"STYLE ID": ids[ids.notna()].astype("int64")})
        out[account] = df.drop_duplicates("STYLE ID").reindex(columns=SHEET_COLUMNS)
        log.info("%s: %d styles from tab %s", account, len(out[account]), title)
    missing = [a for a in accounts if a not in out]
    if len(missing) == len(accounts):
        raise RuntimeError(f"{SOURCE} has no tab of {', '.join(accounts)} (tabs: {', '.join(map(str, tabs))})")
    for a in missing:
        log.warning("%s has no %s tab", SOURCE, a)
    return out


def lookup(ids: list[int]) -> dict[int, dict | str]:
    """Each style's buyer-side answer: straight from www.myntra.com, and through Ultron Phone
    (core/phone.py, the Android app on a mobile line) when myntra.com refuses this machine (403 on
    GitHub's runners) or MYNTRA_ZERO_TD_VIA=phone. The phone answers the same "style" object,
    cut down to what check() reads."""
    from ..core import phone
    from .spider import fetch
    via = (env("MYNTRA_ZERO_TD_VIA") or "").lower()
    answers: dict[int, dict | str] = {}
    if via != "phone":
        answers = fetch(ids, num("MYNTRA_ZERO_TD_CONCURRENCY", 8))
        refused = sum(1 for a in answers.values() if a == "HTTP 403")
        if refused < len(ids) / 2:
            return answers
        log.warning("www.myntra.com refused this machine (%d of %d styles HTTP 403): asking Ultron Phone",
                    refused, len(ids))
    got = phone.run_job("myntra_style", ids, wait_minutes=num("MYNTRA_ZERO_TD_PHONE_MINUTES", 40))
    return {i: got.get(str(i), "no answer from the phone") for i in ids}


def check(answer: dict | str) -> dict:
    """The CHECK_COLUMNS of one style from its buyer-side answer."""
    from .spider import sizes_in_stock
    if not isinstance(answer, dict):
        return {"STATUS": answer if answer == "NOT FOUND" else "NOT CHECKED", "ERROR": answer}
    sizes = sizes_in_stock(answer)
    at_mrp = [s for s in sizes if s["price"] >= s["mrp"] > 0]
    if not sizes:
        status = "OUT OF STOCK"
    elif len(at_mrp) == len(sizes):
        status = "ZERO TD"
    elif at_mrp:
        status = "PART ZERO TD"
    else:
        status = "DISCOUNTED"
    sellers = ", ".join(sorted({str(s.get("displayName") or s.get("sellerName") or "").strip()
                                for s in answer.get("sellers") or []} - {""}))
    return {"STATUS": status,
            "MRP": max((s["mrp"] for s in sizes), default=answer.get("mrp")),
            "SELLING PRICE": min((s["price"] for s in sizes), default=None),
            "SIZES AT MRP": ", ".join(s["size"] for s in at_mrp),
            "STOCK AT MRP": sum(s["stock"] for s in at_mrp),
            "TOTAL STOCK": sum(s["stock"] for s in sizes),
            "SELLER": sellers,
            "LINK": f"https://www.myntra.com/{answer.get('id')}"}


def write_workbook(path: Path, checked: dict[str, pd.DataFrame], when: datetime) -> Path:
    """The workbook sent on Telegram: SUMMARY, ALL ACCOUNTS (every account's zero TD styles together:
    STYLE ID, SKU, SELLING PRICE, each style once), a tab per account (its zero TD styles, most
    stock at MRP first) and NOT CHECKED."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    summary = pd.DataFrame([summarise(a, df) for a, df in checked.items()])
    unchecked = pd.concat([df[df["STATUS"] == "NOT CHECKED"].assign(ACCOUNT=a) for a, df in checked.items()])
    flagged = {a: df[df["STATUS"].isin(FLAGGED)].sort_values("STOCK AT MRP", ascending=False)
               for a, df in checked.items()}
    merged = pd.concat(flagged.values())
    merged = pd.DataFrame({"STYLE ID": merged["STYLE ID"], "SKU": merged["FINAL SKU"].fillna(merged["MY SKU"]),
                           "SELLING PRICE": merged["SELLING PRICE"]}).drop_duplicates("STYLE ID")
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        summary.to_excel(xl, sheet_name="SUMMARY", index=False)
        merged.to_excel(xl, sheet_name=ALL_TAB, index=False)
        for account, df in flagged.items():
            df.reindex(columns=SHEET_COLUMNS + CHECK_COLUMNS).to_excel(xl, sheet_name=account, index=False)
        if len(unchecked):
            unchecked.reindex(columns=["ACCOUNT", "STYLE ID", "MY SKU", "ERROR"]).to_excel(xl, sheet_name="NOT CHECKED", index=False)
        head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="C0392B")
        for ws in xl.book.worksheets:
            for cell in ws[1]:
                cell.font, cell.fill = head, fill
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), 1):
                ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, *(len(str(c.value or "")) + 2 for c in col)))
            if ws.title in checked or ws.title == ALL_TAB:
                link = 1 if ws.title == ALL_TAB else SHEET_COLUMNS.index("STYLE ID") + 1
                for row in ws.iter_rows(min_row=2, min_col=link, max_col=link):
                    row[0].hyperlink = f"https://www.myntra.com/{row[0].value}"
                    row[0].font = Font(color="0563C1", underline="single")
        xl.book["SUMMARY"]["A" + str(len(summary) + 3)] = f"Checked on www.myntra.com at {when:%d %b %Y %I:%M %p} IST"
    return path


def summarise(account: str, df: pd.DataFrame) -> dict:
    status = df["STATUS"].value_counts()
    return {"ACCOUNT": account, "STYLES": len(df), "ZERO TD": int(status.get("ZERO TD", 0)),
            "PART ZERO TD": int(status.get("PART ZERO TD", 0)), "STOCK AT MRP": int(df["STOCK AT MRP"].fillna(0).sum()),
            "DISCOUNTED": int(status.get("DISCOUNTED", 0)), "OUT OF STOCK": int(status.get("OUT OF STOCK", 0)),
            "NOT FOUND": int(status.get("NOT FOUND", 0)), "NOT CHECKED": int(status.get("NOT CHECKED", 0))}


def message(checked: dict[str, pd.DataFrame], when: datetime) -> str:
    rows = [summarise(a, df) for a, df in checked.items()]
    flagged = sum(r["ZERO TD"] + r["PART ZERO TD"] for r in rows)
    out = [f"🏷 <b>Myntra zero TD · {when:%d %b, %I:%M %p}</b>", ""]
    out.append(f"⚠️ {flagged} styles live at MRP" if flagged else "✅ No style live at MRP")
    for r, (account, df) in zip(rows, checked.items()):
        n = r["ZERO TD"] + r["PART ZERO TD"]
        out += ["", f"<b>{account}</b>: {n} at MRP" + (f" ({r['PART ZERO TD']} some sizes)" if r["PART ZERO TD"] else "")
                + (f", {r['STOCK AT MRP']:,} pcs" if n else "")]
        top = df[df["STATUS"].isin(FLAGGED)].sort_values("STOCK AT MRP", ascending=False).head(TOP)
        for _, s in top.iterrows():
            sku = telegram.esc(str(s["FINAL SKU"] if pd.notna(s["FINAL SKU"]) else s["MY SKU"] or "")[:22])
            out.append(f'• <a href="{s["LINK"]}">{s["STYLE ID"]}</a> {sku}: ₹{s["MRP"]:,.0f}, {int(s["STOCK AT MRP"]):,} pcs')
    checked_n = sum(r["STYLES"] - r["NOT CHECKED"] for r in rows)
    unchecked = sum(r["NOT CHECKED"] for r in rows)
    out += ["", f"🔎 {checked_n:,} styles checked on myntra.com"]
    if unchecked:
        out.append(f"⚠️ {unchecked} styles not checked (no answer)")
    if flagged:
        out.append("👉 Put the TD on these styles (sheet attached)")
    return "\n".join(out)


def file_name(when: datetime) -> str:
    """ZERO TD 26-10-26 10.00 AM.xlsx (no ":" - Windows does not allow it in a file name)."""
    return f"ZERO TD {when:%d-%m-%y %I.%M %p}.xlsx"


def save_to_drive(path: Path, when: datetime) -> str:
    """The workbook into AGENT WORK/MYNTRA/MYNTRA ZERO TD SHEETS/OCT-26/26-10-26/; the folder's link."""
    from ..core.drive import XLSX, DriveFolder
    root = DriveFolder(env("AGENT_WORK_FOLDER_ID", AGENT_WORK)).subfolder(AREA, create=True)
    folder = (root.subfolder(FOLDER, create=True).subfolder(f"{when:%b-%y}".upper(), create=True)
              .subfolder(f"{when:%d-%m-%y}", create=True))
    folder.put(path, XLSX)
    return f"https://drive.google.com/drive/folders/{folder.folder_id}"


def run(args) -> int:
    import tempfile
    accounts = [a.upper() for a in (csv_list("MYNTRA_ZERO_TD_ACCOUNTS") or ACCOUNTS)]
    if args.source:
        source = Path(args.source)
    else:
        from ..core.drive import DriveFolder
        source = DriveFolder(env("AGENT_WORK_FOLDER_ID", AGENT_WORK)).subfolder(AREA, create=True).download(
            SOURCE, Path(tempfile.mkdtemp(prefix="zero_td_")) / SOURCE)
        if source is None:
            raise RuntimeError(f"{SOURCE} not found in AGENT WORK/MYNTRA on Drive")
    sheets = read_source(source, accounts)
    if args.limit:
        sheets = {a: df.head(args.limit) for a, df in sheets.items()}

    ids = sorted({int(i) for df in sheets.values() for i in df["STYLE ID"]})
    log.info("Looking up %d styles on www.myntra.com", len(ids))
    answers = lookup(ids)
    found = {i: check(a) for i, a in answers.items()}
    unchecked = sum(1 for c in found.values() if c["STATUS"] == "NOT CHECKED")
    if unchecked == len(ids):
        errors = pd.Series([c["ERROR"] for c in found.values()]).value_counts().head(3)
        raise RuntimeError("www.myntra.com answered no style: " + "; ".join(f"{e} x{n}" for e, n in errors.items()))

    checked = {a: pd.concat([df.reset_index(drop=True),
                             pd.DataFrame([found[int(i)] for i in df["STYLE ID"]])], axis=1)
               for a, df in sheets.items()}
    when = datetime.now()
    path = write_workbook(output_dir("myntra-zero-td") / file_name(when), checked, when)
    msg = message(checked, when)
    if args.no_upload:
        log.info("Not saved to Drive (test run): %s", path)
    else:
        try:
            month = f"{when:%b-%y}".upper()
            msg += f'\n📁 <a href="{save_to_drive(path, when)}">{FOLDER} / {month} / {when:%d-%m-%y}</a>'
        except Exception as exc:  # noqa: BLE001 - the message and the sheet still go out on Telegram
            log.exception("%s not saved to Drive", path.name)
            msg += f"\n⚠️ Not saved to Drive: {telegram.esc(str(exc)[:150])}"
    if args.no_telegram:
        print(msg, f"\n{path}", sep="")
    else:
        telegram.send(msg)
        telegram.send_document(path, f"🏷 Myntra zero TD · {when:%d %b, %I:%M %p}")
    if unchecked > MAX_UNCHECKED * len(ids):
        log.error("%d of %d styles not checked", unchecked, len(ids))
        return 1
    return 0


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="run.py myntra-zero-td",
                                description="Styles live on www.myntra.com with stock at MRP (no TD), per account")
    p.add_argument("--source", help="a local copy of the sheet instead of the one on Drive")
    p.add_argument("--limit", type=int, default=0, help="only the first N styles of each account (a test)")
    p.add_argument("--no-upload", action="store_true", help="keep the workbook in output/, not on Drive")
    p.add_argument("--no-telegram", action="store_true", help="print the message instead of sending it")
    return run_bot("Myntra zero TD", p, run, argv)
